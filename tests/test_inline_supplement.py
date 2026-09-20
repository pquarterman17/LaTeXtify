import shutil
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.shared import Inches
from PIL import Image

from latextify.citations.body_markers import strip_reference_section_to_eof
from latextify.compile.tectonic import compile_document, find_tectonic
from latextify.emit.inline_supplement import (
    BOUNDARY_MARKER,
    mark_inline_supplement,
    render_inline_supplement,
)
from latextify.emit.project import emit_project
from latextify.model import BodyConversionResult


def test_inline_supplement_adds_page_break_and_s_numbering():
    marked = mark_inline_supplement(
        "\\section{Results}\nMain.\n\\section{Supplementary Information}\nSI."
    )
    rendered = render_inline_supplement(marked)

    assert BOUNDARY_MARKER not in rendered
    assert "\\clearpage" in rendered
    assert "\\setcounter{figure}{0}" in rendered
    assert "\\renewcommand{\\thefigure}{S\\arabic{figure}}" in rendered
    assert "\\renewcommand{\\theHfigure}{S.\\arabic{figure}}" in rendered
    assert "\\setcounter{subsection}{0}" in rendered
    assert "\\ifnum\\value{section}=0 S\\arabic{subsection}" in rendered
    assert "\\section*{Supplementary Information}" in rendered
    assert rendered.endswith("SI.")


def test_inline_supplement_consumes_and_reattaches_pandoc_label():
    marked = mark_inline_supplement(
        "\\section{Results}\\label{results}\n"
        "\\section{Supplementary Information}\\label{supplementary-information}\nSI."
    )

    assert "\\section{Supplementary Information}" not in marked
    assert marked.count("\\label{supplementary-information}") == 1
    assert (
        "%%LATEXTIFY_INLINE_SUPPLEMENT%%\n"
        "\\section*{Supplementary Information}\\label{supplementary-information}"
    ) in marked


def test_main_references_are_removed_without_removing_inline_supplement():
    marked = mark_inline_supplement(
        "Main.\n\\section{References}\nReference 1.\n\\section{Supplemental Material}\nSI content."
    )

    stripped = strip_reference_section_to_eof(marked)

    assert "Reference 1" not in stripped
    assert BOUNDARY_MARKER in stripped
    assert "SI content" in stripped


def test_both_main_and_supplement_reference_lists_are_removed():
    marked = mark_inline_supplement(
        "Main.\n\\section{References}\nMain ref.\n"
        "\\section{Supplementary Material}\nSI.\n"
        "\\section{References}\nSI ref."
    )

    stripped = strip_reference_section_to_eof(marked)

    assert "Main ref" not in stripped
    assert "SI ref" not in stripped
    assert "SI." in stripped


def test_inline_supplement_accepts_a_bold_plain_heading():
    marked = mark_inline_supplement("Main.\n\n\\textbf{Supplementary Information}\n\nSI.")
    assert BOUNDARY_MARKER in marked


def test_inline_supplement_requires_a_recognized_heading():
    try:
        mark_inline_supplement("\\section{Results}\nNo supplement here.")
    except ValueError as exc:
        assert "no heading named Supplementary" in str(exc)
    else:
        raise AssertionError("missing supplement heading should be rejected")


def test_emit_project_writes_inline_supplement_into_main_body(tmp_path, monkeypatch):
    source = Path(__file__).parent / "fixtures" / "clean.docx"
    docx = tmp_path / "merged.docx"
    shutil.copy(source, docx)

    def fake_convert(_docx_path, media_dir, **_kwargs):
        return BodyConversionResult(
            tex="\\section{Results}\nMain.\n\\section{Supplementary Material}\nSI.",
            media_dir=media_dir,
            figure_count=0,
            citation_count=0,
        )

    monkeypatch.setattr("latextify.emit.project.convert_docx_to_body", fake_convert)

    result = emit_project(docx, "revtex4-2", tmp_path / "output", inline_supplement=True)
    body = result.body_tex_path.read_text(encoding="utf-8")

    assert result.supplement is None
    assert not (result.output_dir / "supplement.tex").exists()
    assert "\\clearpage" in body
    assert "\\section*{Supplementary Material}" in body


def test_inline_supplement_rejects_figures_at_end_before_writing(tmp_path):
    source = Path(__file__).parent / "fixtures" / "clean.docx"
    docx = tmp_path / "merged.docx"
    shutil.copy(source, docx)
    output = tmp_path / "output"

    with pytest.raises(ValueError, match="figures-at-end"):
        emit_project(
            docx,
            "revtex4-2",
            output,
            inline_supplement=True,
            figures_at_end=True,
        )

    assert not output.exists()


def test_missing_inline_heading_does_not_modify_existing_output(tmp_path, monkeypatch):
    source = Path(__file__).parent / "fixtures" / "clean.docx"
    docx = tmp_path / "merged.docx"
    shutil.copy(source, docx)
    figures = tmp_path / "output" / "revtex4-2" / "figures"
    figures.mkdir(parents=True)
    existing = figures / "fig1.png"
    existing.write_bytes(b"existing figure")

    def fake_convert(_docx_path, media_dir, **_kwargs):
        return BodyConversionResult(
            tex="\\section{Results}\nNo supplement heading.",
            media_dir=media_dir,
            figure_count=0,
            citation_count=0,
        )

    monkeypatch.setattr("latextify.emit.project.convert_docx_to_body", fake_convert)

    with pytest.raises(ValueError, match="no heading named Supplementary"):
        emit_project(docx, "revtex4-2", tmp_path / "output", inline_supplement=True)

    assert existing.read_bytes() == b"existing figure"


def _write_windows_emf(path: Path) -> None:
    """Create a real, tiny EMF through Windows GDI for the end-to-end test."""
    import ctypes
    from ctypes import wintypes

    gdi32 = ctypes.WinDLL("gdi32", use_last_error=True)
    gdi32.CreateEnhMetaFileW.argtypes = (
        wintypes.HDC,
        wintypes.LPCWSTR,
        ctypes.POINTER(wintypes.RECT),
        wintypes.LPCWSTR,
    )
    gdi32.CreateEnhMetaFileW.restype = wintypes.HANDLE
    gdi32.Rectangle.argtypes = (
        wintypes.HDC,
        ctypes.c_int,
        ctypes.c_int,
        ctypes.c_int,
        ctypes.c_int,
    )
    gdi32.CloseEnhMetaFile.argtypes = (wintypes.HDC,)
    gdi32.CloseEnhMetaFile.restype = wintypes.HANDLE
    gdi32.DeleteEnhMetaFile.argtypes = (wintypes.HANDLE,)
    frame = wintypes.RECT(0, 0, 2540, 1270)
    hdc = gdi32.CreateEnhMetaFileW(None, str(path), ctypes.byref(frame), None)
    if not hdc:
        raise ctypes.WinError(ctypes.get_last_error())
    gdi32.Rectangle(hdc, 100, 100, 2400, 1100)
    metafile = gdi32.CloseEnhMetaFile(hdc)
    if not metafile:
        raise ctypes.WinError(ctypes.get_last_error())
    gdi32.DeleteEnhMetaFile(metafile)


def _replace_docx_pngs_with_emfs(docx_path: Path, emfs: tuple[Path, ...]) -> None:
    rebuilt = docx_path.with_suffix(".rebuilt.docx")
    with (
        zipfile.ZipFile(docx_path) as source,
        zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as destination,
    ):
        for info in source.infolist():
            if info.filename in {f"word/media/image{index}.png" for index in range(1, 3)}:
                continue
            data = source.read(info.filename)
            if info.filename == "word/_rels/document.xml.rels":
                data = data.replace(b".png", b".emf")
            elif info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Default Extension="emf" ContentType="image/x-emf"/></Types>',
                )
            destination.writestr(info, data)
        for index, emf in enumerate(emfs, start=1):
            destination.writestr(f"word/media/image{index}.emf", emf.read_bytes())
    rebuilt.replace(docx_path)


@pytest.mark.skipif(sys.platform != "win32", reason="real EMF creation uses Windows GDI")
def test_real_merged_docx_with_embedded_emf_and_column_choices(tmp_path):
    first = tmp_path / "main.emf"
    second = tmp_path / "supplement.emf"
    _write_windows_emf(first)
    _write_windows_emf(second)
    first_png = tmp_path / "main.png"
    second_png = tmp_path / "supplement.png"
    Image.new("RGB", (200, 100), "white").save(first_png)
    Image.new("RGB", (200, 100), "white").save(second_png)
    doc = Document()
    doc.add_heading("Results", level=1)
    doc.add_picture(str(first_png), width=Inches(2))
    doc.add_paragraph("Figure 1. Main figure")
    doc.add_heading("Supplementary Information", level=1)
    doc.add_picture(str(second_png), width=Inches(2))
    doc.add_paragraph("Figure S1. Supplementary figure")
    merged = tmp_path / "merged-emf.docx"
    doc.save(merged)
    _replace_docx_pngs_with_emfs(merged, (first, second))

    result = emit_project(
        merged,
        "revtex4-2",
        tmp_path / "output",
        inline_supplement=True,
        figure_placements={("", 1): "one", ("", 2): "two"},
    )
    body = result.body_tex_path.read_text(encoding="utf-8")

    assert result.figure_count == 2
    assert "\\begin{figure}\n" in body
    assert "\\begin{figure*}\n" in body
    assert "\\clearpage" in body
    assert len(tuple(result.figures_dir.glob("fig*.*"))) == 2
    tectonic = find_tectonic()
    if tectonic is not None:
        compiled = compile_document(result.main_tex_path, tectonic_path=tectonic)
        assert compiled.success, compiled.raw_log

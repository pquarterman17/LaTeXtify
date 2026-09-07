"""``latextify figures``, ``convert --vector-figures``, and the report's figure line.

The user-facing half of the figure inventory (the classification itself is
covered by ``tests/test_figures_inventory.py``). What matters here is that an
author can answer two questions without opening the code:

    * which figure is number 3, so an override file can be NAMED correctly, and
    * after converting, which figures are still screenshots.

Fixtures are built at run time with python-docx, Pillow and reportlab.
"""

from __future__ import annotations

import io
import json
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image
from reportlab.pdfgen import canvas
from typer.testing import CliRunner

from latextify.cli import app
from latextify.emit.project import emit_project

runner = CliRunner()


def _png_bytes(size: tuple[int, int]) -> io.BytesIO:
    buf = io.BytesIO()
    Image.new("RGB", size, (40, 90, 160)).save(buf, format="PNG")
    buf.seek(0)
    return buf


def _vector_pdf(path: Path, size: tuple[float, float] = (400, 200)) -> Path:
    c = canvas.Canvas(str(path), pagesize=size)
    c.drawString(20, size[1] / 2, "Wavelength (nm)")
    for i in range(400):
        c.line(i % int(size[0]), 0, i % int(size[0]), 40)
    c.save()
    return path


def _manuscript(path: Path, *, sizes=((400, 300), (400, 300))) -> Path:
    """A manuscript with one captioned figure per entry in ``sizes``."""
    document = Document()
    document.add_heading("Figure Inventory Fixture", level=0)
    document.add_heading("Results", level=1)
    for index, size in enumerate(sizes, start=1):
        document.add_paragraph(f"Result {index}, illustrated below.")
        document.add_picture(_png_bytes(size), width=Inches(3))
        document.add_paragraph(f"Figure {index}: Result number {index}.", style="Caption")
    document.save(path)
    return path


# --------------------------------------------------------------------------- #
# latextify figures -- the listing
# --------------------------------------------------------------------------- #


def test_listing_shows_each_figure_with_its_number_and_caption(tmp_path):
    result = runner.invoke(app, ["figures", str(_manuscript(tmp_path / "paper.docx"))])

    assert result.exit_code == 0, result.output
    assert "2 figure(s)" in result.output
    # The number is what names an override file, so it must be shown per row.
    assert "Result number 1." in result.output
    assert "Result number 2." in result.output
    assert "raster" in result.output


def test_listing_flags_low_resolution_and_names_the_remedy(tmp_path):
    result = runner.invoke(app, ["figures", str(_manuscript(tmp_path / "paper.docx"))])

    assert "LOW" in result.output
    assert "0 of 2 already vector." in result.output
    assert "figures/fig<N>.pdf" in result.output


def test_listing_reports_a_supplied_override_as_vector(tmp_path):
    """The confirmation loop: supply a vector file, see it recognised."""
    docx = _manuscript(tmp_path / "paper.docx")
    (tmp_path / "figures").mkdir()
    _vector_pdf(tmp_path / "figures" / "fig1.pdf")

    result = runner.invoke(app, ["figures", str(docx)])

    assert result.exit_code == 0, result.output
    assert "override" in result.output
    assert "vector" in result.output
    assert "1 of 2 already vector." in result.output
    # Figure 1 is handled; only figure 2 is still worth replacing.
    assert "Figure(s) 2 would print below" in result.output


def test_listing_writes_nothing_into_the_manuscript_directory(tmp_path):
    docx = _manuscript(tmp_path / "paper.docx")
    before = sorted(p.name for p in tmp_path.iterdir())

    assert runner.invoke(app, ["figures", str(docx)]).exit_code == 0

    assert sorted(p.name for p in tmp_path.iterdir()) == before


def test_listing_as_json_is_machine_readable(tmp_path):
    docx = _manuscript(tmp_path / "paper.docx", sizes=((400, 300),))
    result = runner.invoke(app, ["figures", str(docx), "--json"])

    assert result.exit_code == 0, result.output
    body = json.loads(result.output)
    (figure,) = body["figures"]
    assert figure["number"] == 1
    assert figure["kind"] == "raster"
    assert figure["needs_attention"] is True
    assert body["caption_gaps"] == []


def test_listing_warns_when_a_caption_has_no_image_behind_it(tmp_path):
    """A gap shifts every later figure's number, which is exactly how an
    override lands on the wrong figure -- so the listing has to say so."""
    document = Document()
    document.add_heading("Gap Fixture", level=0)
    document.add_heading("Results", level=1)
    for number in (1, 2, 3):
        document.add_paragraph(f"Result {number}.")
        if number != 2:  # figure 2 is captioned but never pasted in
            document.add_picture(_png_bytes((400, 300)), width=Inches(3))
        document.add_paragraph(f"Figure {number}: Figure {number} text.", style="Caption")
    docx = tmp_path / "gap.docx"
    document.save(docx)

    result = runner.invoke(app, ["figures", str(docx)])

    assert result.exit_code == 0, result.output
    assert "Caption(s) with no image" in result.output


def test_listing_a_document_with_no_figures_says_so(tmp_path):
    document = Document()
    document.add_heading("No Figures", level=0)
    document.add_paragraph("Text only.")
    docx = tmp_path / "plain.docx"
    document.save(docx)

    result = runner.invoke(app, ["figures", str(docx)])

    assert result.exit_code == 0, result.output
    assert "no figures found" in result.output


def test_listing_a_missing_file_is_a_usage_error(tmp_path):
    result = runner.invoke(app, ["figures", str(tmp_path / "absent.docx")])
    assert result.exit_code != 0


# --------------------------------------------------------------------------- #
# convert --vector-figures -- the per-figure warning
# --------------------------------------------------------------------------- #


def test_vector_figures_warns_once_per_raster_figure(tmp_path):
    docx = _manuscript(tmp_path / "paper.docx")

    result = runner.invoke(
        app,
        ["convert", str(docx), "-j", "revtex4-2", "-o", str(tmp_path / "out"), "--vector-figures"],
    )

    assert result.exit_code == 0, result.output
    assert "figure 1 is a raster image" in result.output
    assert "figure 2 is a raster image" in result.output
    assert "figures/fig1.pdf" in result.output


def test_vector_figures_is_silent_about_a_supplied_vector_override(tmp_path):
    docx = _manuscript(tmp_path / "paper.docx")
    (tmp_path / "figures").mkdir()
    _vector_pdf(tmp_path / "figures" / "fig1.pdf")

    result = runner.invoke(
        app,
        ["convert", str(docx), "-j", "revtex4-2", "-o", str(tmp_path / "out"), "--vector-figures"],
    )

    assert result.exit_code == 0, result.output
    assert "figure 1 is a raster" not in result.output
    assert "figure 2 is a raster image" in result.output


def test_without_the_flag_no_vector_advice_is_emitted(tmp_path):
    """Off by default: an existing conversion's warnings must not change shape."""
    docx = _manuscript(tmp_path / "paper.docx")

    result = runner.invoke(
        app, ["convert", str(docx), "-j", "revtex4-2", "-o", str(tmp_path / "out")]
    )

    assert result.exit_code == 0, result.output
    assert "raster image" not in result.output


def test_a_converted_svg_is_not_reported_as_needing_replacement(tmp_path):
    """The emitter describes what it WROTE.

    An SVG override is converted to PDF on the way into figures/. Reporting the
    source would tell an author to replace a figure that is already vector in
    what actually ships.
    """
    docx = _manuscript(tmp_path / "paper.docx", sizes=((400, 300),))
    (tmp_path / "figures").mkdir()
    (tmp_path / "figures" / "fig1.svg").write_text(
        '<svg xmlns="http://www.w3.org/2000/svg" width="400" height="200">'
        '<rect width="400" height="200" fill="#3050a0"/></svg>',
        encoding="utf-8",
    )

    result = runner.invoke(
        app,
        ["convert", str(docx), "-j", "revtex4-2", "-o", str(tmp_path / "out"), "--vector-figures"],
    )

    assert result.exit_code == 0, result.output
    assert "figure 1 is a raster" not in result.output


# --------------------------------------------------------------------------- #
# report.md -- the durable record of what shipped
# --------------------------------------------------------------------------- #


def test_report_records_which_figures_are_vector(tmp_path):
    docx = _manuscript(tmp_path / "paper.docx")
    (tmp_path / "figures").mkdir()
    _vector_pdf(tmp_path / "figures" / "fig1.pdf")

    result = emit_project(docx, "revtex4-2", tmp_path / "out")
    report = result.report_path.read_text(encoding="utf-8")

    assert "**Fig 1** (OVERRIDE, vector)" in report
    assert "**Fig 2** (EMBEDDED, raster" in report
    assert "below 300 DPI" in report


def test_report_figure_line_survives_an_unmeasurable_figure(tmp_path):
    """A figure the inventory cannot describe reads exactly as it always did."""
    docx = _manuscript(tmp_path / "paper.docx", sizes=((400, 300),))
    result = emit_project(docx, "revtex4-2", tmp_path / "out")
    for written in (result.figures_dir).glob("fig1.*"):
        written.write_bytes(b"corrupted after the fact")

    from latextify.report.render import write_report

    write_report(result.report_path, preflight=None, emit_result=result, reconciliation=None)
    report = result.report_path.read_text(encoding="utf-8")

    assert "**Fig 1** (EMBEDDED)" in report


# --------------------------------------------------------------------------- #
# Findings from the self-review: JSON paths, threshold rounding, exit code
# --------------------------------------------------------------------------- #


def test_json_never_reports_a_path_into_a_deleted_temporary_directory(tmp_path):
    """--json is documented as scriptable, so its paths must be usable.

    An embedded figure's media is extracted into a TemporaryDirectory already
    gone by the time the output prints; reporting it would be a dangling path.
    """
    docx = _manuscript(tmp_path / "paper.docx", sizes=((400, 300),))
    (tmp_path / "figures").mkdir()
    override = _vector_pdf(tmp_path / "figures" / "fig1.pdf")

    embedded = json.loads(runner.invoke(app, ["figures", str(docx), "--json"]).output)
    assert embedded["figures"][0]["source"] == "override"
    assert embedded["figures"][0]["path"] == str(override)

    override.unlink()  # now nothing is supplied, so the figure is embedded
    body = json.loads(runner.invoke(app, ["figures", str(docx), "--json"]).output)
    assert body["figures"][0]["source"] == "embedded"
    assert body["figures"][0]["path"] is None


def test_the_low_flag_and_the_summary_never_disagree(tmp_path):
    """Rounding for display before comparing showed 299.7 DPI as an unflagged
    "300" while the summary below still listed it as needing replacement."""
    from latextify.cli_figures import _dpi_cell
    from latextify.figures.inventory import MIN_PRINT_DPI, describe
    from latextify.model.figure import Figure, FigureSource

    # 1019 px over 3.4 in is 299.7 DPI: rounds to 300, is below the floor.
    png = tmp_path / "borderline.png"
    Image.new("RGB", (1019, 1019), (40, 90, 160)).save(png)
    facts = describe(Figure(number=1, caption="c", embedded_path=png, source=FigureSource.EMBEDDED))

    assert facts.dpi < MIN_PRINT_DPI
    assert facts.needs_attention is True
    assert "LOW" in _dpi_cell(facts)  # the row agrees with the summary


def test_an_unreadable_manuscript_exits_one_like_every_other_command(tmp_path):
    """convert/export/equations all exit 1 for the same ValueError."""
    bogus = tmp_path / "not-really.docx"
    bogus.write_bytes(b"certainly not a docx")

    result = runner.invoke(app, ["figures", str(bogus)])

    assert result.exit_code == 1


# --------------------------------------------------------------------------- #
# Supplementary material gets the same treatment as the main document
# --------------------------------------------------------------------------- #


def test_vector_figures_reports_supplement_figures_too(tmp_path):
    """The flag says "every figure"; SI figures were silently exempt because
    emit_supplement copied figures on its own path instead of the shared one."""
    main = _manuscript(tmp_path / "paper.docx", sizes=((400, 300),))
    si = _manuscript(tmp_path / "si.docx", sizes=((400, 300),))

    result = runner.invoke(
        app,
        [
            "convert",
            str(main),
            "-j",
            "revtex4-2",
            "-o",
            str(tmp_path / "out"),
            "--supplement",
            str(si),
            "--vector-figures",
        ],
    )

    assert result.exit_code == 0, result.output
    # The SI figure is named figS1, and its warning is supplement-scoped.
    assert "figures/figS1.pdf" in result.output
    assert "supplement:" in result.output


def test_a_supplied_supplement_vector_override_is_not_reported(tmp_path):
    main = _manuscript(tmp_path / "paper.docx", sizes=((400, 300),))
    si = _manuscript(tmp_path / "si.docx", sizes=((400, 300),))
    (tmp_path / "figures").mkdir()
    _vector_pdf(tmp_path / "figures" / "figS1.pdf")

    result = runner.invoke(
        app,
        [
            "convert",
            str(main),
            "-j",
            "revtex4-2",
            "-o",
            str(tmp_path / "out"),
            "--supplement",
            str(si),
            "--vector-figures",
        ],
    )

    assert result.exit_code == 0, result.output
    assert "figures/figS1.pdf" not in result.output
    assert "figures/fig1.pdf" in result.output  # the main document still is

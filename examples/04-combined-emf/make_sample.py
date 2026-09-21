"""Create the Windows acceptance sample: merged main/SI DOCX with real EMFs."""

from __future__ import annotations

import ctypes
import zipfile
from ctypes import wintypes
from pathlib import Path

from docx import Document
from docx.shared import Inches
from PIL import Image

HERE = Path(__file__).resolve().parent
OUTPUT = HERE / "Combined-Manuscript-with-EMF.docx"


def _write_emf(path: Path, *, tall: bool) -> None:
    gdi32 = ctypes.WinDLL("gdi32", use_last_error=True)
    gdi32.CreateEnhMetaFileW.argtypes = (
        wintypes.HDC,
        wintypes.LPCWSTR,
        ctypes.POINTER(wintypes.RECT),
        wintypes.LPCWSTR,
    )
    gdi32.CreateEnhMetaFileW.restype = wintypes.HANDLE
    gdi32.Rectangle.argtypes = (
        wintypes.HDC,
        ctypes.c_int,
        ctypes.c_int,
        ctypes.c_int,
        ctypes.c_int,
    )
    gdi32.CloseEnhMetaFile.argtypes = (wintypes.HDC,)
    gdi32.CloseEnhMetaFile.restype = wintypes.HANDLE
    gdi32.DeleteEnhMetaFile.argtypes = (wintypes.HANDLE,)
    frame = wintypes.RECT(0, 0, 1270 if tall else 2540, 2540 if tall else 1270)
    hdc = gdi32.CreateEnhMetaFileW(None, str(path), ctypes.byref(frame), "LaTeXtify\0EMF sample\0")
    if not hdc:
        raise ctypes.WinError(ctypes.get_last_error())
    gdi32.Rectangle(hdc, 100, 100, frame.right - 100, frame.bottom - 100)
    metafile = gdi32.CloseEnhMetaFile(hdc)
    if not metafile:
        raise ctypes.WinError(ctypes.get_last_error())
    gdi32.DeleteEnhMetaFile(metafile)


def _replace_pngs(docx_path: Path, emfs: tuple[Path, ...]) -> None:
    rebuilt = docx_path.with_suffix(".rebuilt.docx")
    with (
        zipfile.ZipFile(docx_path) as source,
        zipfile.ZipFile(rebuilt, "w", zipfile.ZIP_DEFLATED) as destination,
    ):
        image_names = {f"word/media/image{index}.png" for index in range(1, len(emfs) + 1)}
        for info in source.infolist():
            if info.filename in image_names:
                continue
            data = source.read(info.filename)
            if info.filename == "word/_rels/document.xml.rels":
                data = data.replace(b".png", b".emf")
            elif info.filename == "[Content_Types].xml":
                data = data.replace(
                    b"</Types>",
                    b'<Default Extension="emf" ContentType="image/x-emf"/></Types>',
                )
            destination.writestr(info, data)
        for index, emf in enumerate(emfs, start=1):
            destination.writestr(f"word/media/image{index}.emf", emf.read_bytes())
    rebuilt.replace(docx_path)


def main() -> None:
    if not hasattr(ctypes, "WinDLL"):
        raise SystemExit("This sample must be generated on Windows.")
    main_emf, supplement_emf = HERE / "main.emf", HERE / "supplement.emf"
    main_png, supplement_png = HERE / "main.png", HERE / "supplement.png"
    try:
        _write_emf(main_emf, tall=True)
        _write_emf(supplement_emf, tall=False)
        Image.new("RGB", (200, 400), "white").save(main_png)
        Image.new("RGB", (400, 200), "white").save(supplement_png)

        doc = Document()
        doc.core_properties.title = "LaTeXtify combined manuscript acceptance sample"
        doc.core_properties.author = "Example Author"
        doc.add_heading("A Combined Manuscript Acceptance Sample", level=0)
        doc.add_paragraph("Example Author")
        doc.add_heading("Abstract", level=1)
        doc.add_paragraph(
            "This synthetic document verifies offline conversion without containing research data."
        )
        doc.add_heading("Results", level=1)
        doc.add_paragraph("The main text contains an embedded Windows Enhanced Metafile.")
        doc.add_picture(str(main_png), width=Inches(2.0))
        doc.add_paragraph("Figure 1. Main-text EMF intended for one column.")
        doc.add_paragraph("A demonstration result is cited here [1].")
        doc.add_heading("References", level=1)
        doc.add_paragraph("[1] Example, A. A synthetic reference. Example Journal 1, 1–2 (2026).")
        doc.add_heading("Supplementary Information", level=1)
        doc.add_paragraph("This section must begin on a new page and use S-numbering.")
        doc.add_picture(str(supplement_png), width=Inches(5.5))
        doc.add_paragraph("Figure S1. Supplementary EMF intended to span two columns.")
        doc.add_heading("Supplementary Methods", level=2)
        doc.add_paragraph("No external data or network service is required for this sample.")
        doc.save(OUTPUT)
        _replace_pngs(OUTPUT, (main_emf, supplement_emf))
        print(OUTPUT)
    finally:
        for temporary in (main_emf, supplement_emf, main_png, supplement_png):
            temporary.unlink(missing_ok=True)


if __name__ == "__main__":
    main()
